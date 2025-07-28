"""
PDF Processing module for Sustainable Agent Consultant
Handles extraction and preprocessing of text from PDF sustainability reports
"""

import os
import re
import logging
from typing import List, Dict, Optional, Tuple
from pathlib import Path

import PyPDF2
import pdfplumber
from pypdf import PdfReader
import pandas as pd
from docx import Document

from config import Config

# Set up logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

class PDFProcessor:
    """Handles PDF text extraction and preprocessing for sustainability reports"""
    
    def __init__(self):
        self.config = Config()
        self.supported_formats = self.config.SUPPORTED_FORMATS
        
    def validate_file(self, file_path: str) -> bool:
        """Validate if the file is supported and within size limits"""
        try:
            path = Path(file_path)
            
            # Check if file exists
            if not path.exists():
                logger.error(f"File does not exist: {file_path}")
                return False
            
            # Check file extension
            if path.suffix.lower() not in self.supported_formats:
                logger.error(f"Unsupported file format: {path.suffix}")
                return False
            
            # Check file size
            file_size_mb = path.stat().st_size / (1024 * 1024)
            if file_size_mb > self.config.MAX_FILE_SIZE_MB:
                logger.error(f"File too large: {file_size_mb}MB > {self.config.MAX_FILE_SIZE_MB}MB")
                return False
            
            return True
            
        except Exception as e:
            logger.error(f"Error validating file: {str(e)}")
            return False
    
    def extract_text_from_pdf(self, pdf_path: str, method: str = "pdfplumber") -> str:
        """Extract text from PDF using different methods"""
        try:
            if method == "pdfplumber":
                return self._extract_with_pdfplumber(pdf_path)
            elif method == "pypdf2":
                return self._extract_with_pypdf2(pdf_path)
            elif method == "pypdf":
                return self._extract_with_pypdf(pdf_path)
            else:
                # Try all methods and return the one with most text
                texts = []
                for m in ["pdfplumber", "pypdf", "pypdf2"]:
                    try:
                        text = self._extract_text_by_method(pdf_path, m)
                        texts.append((m, text))
                    except:
                        continue
                
                if texts:
                    # Return the method that extracted the most text
                    best_method, best_text = max(texts, key=lambda x: len(x[1]))
                    logger.info(f"Best extraction method: {best_method}")
                    return best_text
                else:
                    raise Exception("All extraction methods failed")
                    
        except Exception as e:
            logger.error(f"Error extracting text from PDF: {str(e)}")
            return ""
    
    def _extract_with_pdfplumber(self, pdf_path: str) -> str:
        """Extract text using pdfplumber (best for tables and complex layouts)"""
        text = ""
        with pdfplumber.open(pdf_path) as pdf:
            for page_num, page in enumerate(pdf.pages):
                try:
                    page_text = page.extract_text()
                    if page_text:
                        text += f"\n--- Page {page_num + 1} ---\n"
                        text += page_text
                        
                    # Extract tables if any
                    tables = page.extract_tables()
                    for table_num, table in enumerate(tables):
                        if table:
                            text += f"\n--- Table {table_num + 1} on Page {page_num + 1} ---\n"
                            for row in table:
                                if row:
                                    text += " | ".join([cell or "" for cell in row]) + "\n"
                                    
                except Exception as e:
                    logger.warning(f"Error processing page {page_num + 1}: {str(e)}")
                    continue
        
        return text
    
    def _extract_with_pypdf2(self, pdf_path: str) -> str:
        """Extract text using PyPDF2"""
        text = ""
        with open(pdf_path, 'rb') as file:
            pdf_reader = PyPDF2.PdfReader(file)
            for page_num, page in enumerate(pdf_reader.pages):
                try:
                    page_text = page.extract_text()
                    if page_text:
                        text += f"\n--- Page {page_num + 1} ---\n"
                        text += page_text
                except Exception as e:
                    logger.warning(f"Error processing page {page_num + 1}: {str(e)}")
                    continue
        
        return text
    
    def _extract_with_pypdf(self, pdf_path: str) -> str:
        """Extract text using pypdf"""
        text = ""
        reader = PdfReader(pdf_path)
        for page_num, page in enumerate(reader.pages):
            try:
                page_text = page.extract_text()
                if page_text:
                    text += f"\n--- Page {page_num + 1} ---\n"
                    text += page_text
            except Exception as e:
                logger.warning(f"Error processing page {page_num + 1}: {str(e)}")
                continue
        
        return text
    
    def _extract_text_by_method(self, pdf_path: str, method: str) -> str:
        """Helper method to extract text by specific method"""
        if method == "pdfplumber":
            return self._extract_with_pdfplumber(pdf_path)
        elif method == "pypdf2":
            return self._extract_with_pypdf2(pdf_path)
        elif method == "pypdf":
            return self._extract_with_pypdf(pdf_path)
        else:
            raise ValueError(f"Unknown extraction method: {method}")
    
    def extract_text_from_docx(self, docx_path: str) -> str:
        """Extract text from DOCX files"""
        try:
            doc = Document(docx_path)
            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            
            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text for cell in row.cells])
                    text += row_text + "\n"
            
            return text
            
        except Exception as e:
            logger.error(f"Error extracting text from DOCX: {str(e)}")
            return ""
    
    def extract_text_from_txt(self, txt_path: str) -> str:
        """Extract text from TXT files"""
        try:
            with open(txt_path, 'r', encoding='utf-8') as file:
                return file.read()
        except UnicodeDecodeError:
            # Try with different encoding
            with open(txt_path, 'r', encoding='latin-1') as file:
                return file.read()
        except Exception as e:
            logger.error(f"Error extracting text from TXT: {str(e)}")
            return ""
    
    def preprocess_text(self, text: str) -> str:
        """Clean and preprocess extracted text"""
        if not text:
            return ""
        
        # Remove extra whitespace and normalize
        text = re.sub(r'\s+', ' ', text)
        
        # Remove special characters but keep important punctuation
        text = re.sub(r'[^\w\s.,;:!?()-]', ' ', text)
        
        # Remove page breaks and headers/footers patterns
        text = re.sub(r'--- Page \d+ ---', '\n', text)
        text = re.sub(r'Page \d+ of \d+', '', text)
        
        # Remove excessive newlines
        text = re.sub(r'\n\s*\n', '\n\n', text)
        
        return text.strip()
    
    def chunk_text(self, text: str, chunk_size: int = None, overlap: int = None) -> List[str]:
        """Split text into chunks for embedding"""
        chunk_size = chunk_size or self.config.CHUNK_SIZE
        overlap = overlap or self.config.CHUNK_OVERLAP
        
        if not text:
            return []
        
        chunks = []
        start = 0
        text_length = len(text)
        
        while start < text_length:
            # Find the end of this chunk
            end = start + chunk_size
            
            # If we're not at the end of the text, try to break at a sentence or paragraph
            if end < text_length:
                # Look for sentence endings near the chunk boundary
                sentence_ends = [text.rfind('.', start, end), 
                               text.rfind('!', start, end), 
                               text.rfind('?', start, end)]
                
                valid_ends = [e for e in sentence_ends if e > start + chunk_size * 0.8]
                if valid_ends:
                    best_end = max(valid_ends)
                    if best_end > start:
                        end = best_end + 1
            
            chunk = text[start:end].strip()
            if chunk:
                chunks.append(chunk)
            
            # Move start position (with overlap)
            start = max(start + chunk_size - overlap, end - overlap)
            
            # Avoid infinite loops
            if start >= text_length:
                break
        
        return chunks
    
    def extract_sustainability_metrics(self, text: str) -> Dict[str, List[str]]:
        """Extract sustainability-related metrics and data from text"""
        metrics = {
            'carbon_emissions': [],
            'energy_consumption': [],
            'water_usage': [],
            'waste_generation': [],
            'renewable_energy': [],
            'targets': [],
            'certifications': []
        }
        
        # Patterns for different metrics
        patterns = {
            'carbon_emissions': [
                r'(\d+(?:,\d+)*(?:\.\d+)?)\s*(?:tons?|tonnes?|kg|metric tons?)\s*(?:of\s+)?(?:CO2|carbon|emissions?)',
                r'carbon footprint.*?(\d+(?:,\d+)*(?:\.\d+)?)',
                r'GHG emissions.*?(\d+(?:,\d+)*(?:\.\d+)?)'
            ],
            'energy_consumption': [
                r'(\d+(?:,\d+)*(?:\.\d+)?)\s*(?:MWh|GWh|kWh|TWh)',
                r'energy consumption.*?(\d+(?:,\d+)*(?:\.\d+)?)',
                r'electricity.*?(\d+(?:,\d+)*(?:\.\d+)?)\s*(?:MWh|GWh|kWh)'
            ],
            'water_usage': [
                r'(\d+(?:,\d+)*(?:\.\d+)?)\s*(?:liters?|litres?|gallons?|m3|cubic meters?)',
                r'water consumption.*?(\d+(?:,\d+)*(?:\.\d+)?)',
                r'water usage.*?(\d+(?:,\d+)*(?:\.\d+)?)'
            ],
            'renewable_energy': [
                r'(\d+(?:,\d+)*(?:\.\d+)?)\s*%.*?renewable',
                r'renewable energy.*?(\d+(?:,\d+)*(?:\.\d+)?)',
                r'solar.*?(\d+(?:,\d+)*(?:\.\d+)?)\s*(?:MW|GW|kW)'
            ]
        }
        
        for metric_type, metric_patterns in patterns.items():
            for pattern in metric_patterns:
                matches = re.finditer(pattern, text, re.IGNORECASE)
                for match in matches:
                    context_start = max(0, match.start() - 100)
                    context_end = min(len(text), match.end() + 100)
                    context = text[context_start:context_end].strip()
                    metrics[metric_type].append(context)
        
        return metrics
    
    def process_file(self, file_path: str) -> Dict[str, any]:
        """Main method to process a file and return extracted information"""
        if not self.validate_file(file_path):
            return {"error": "File validation failed"}
        
        file_path = Path(file_path)
        extension = file_path.suffix.lower()
        
        try:
            # Extract text based on file type
            if extension == '.pdf':
                raw_text = self.extract_text_from_pdf(str(file_path))
            elif extension == '.docx':
                raw_text = self.extract_text_from_docx(str(file_path))
            elif extension == '.txt':
                raw_text = self.extract_text_from_txt(str(file_path))
            else:
                return {"error": f"Unsupported file format: {extension}"}
            
            if not raw_text:
                return {"error": "No text could be extracted from the file"}
            
            # Preprocess text
            processed_text = self.preprocess_text(raw_text)
            
            # Create chunks
            chunks = self.chunk_text(processed_text)
            
            # Extract sustainability metrics
            metrics = self.extract_sustainability_metrics(processed_text)
            
            return {
                "filename": file_path.name,
                "raw_text": raw_text,
                "processed_text": processed_text,
                "chunks": chunks,
                "metrics": metrics,
                "chunk_count": len(chunks),
                "text_length": len(processed_text),
                "status": "success"
            }
            
        except Exception as e:
            logger.error(f"Error processing file {file_path}: {str(e)}")
            return {"error": f"Processing failed: {str(e)}"}


# Example usage and testing
if __name__ == "__main__":
    processor = PDFProcessor()
    
    # Test with a sample file (you would replace this with actual file path)
    # result = processor.process_file("sample_sustainability_report.pdf")
    # print(f"Processing result: {result.get('status', 'failed')}")
    # print(f"Number of chunks: {result.get('chunk_count', 0)}")
    # print(f"Text length: {result.get('text_length', 0)}")
    
    print("PDF Processor module loaded successfully!")